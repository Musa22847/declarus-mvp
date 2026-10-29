"""Экспорт ответа модели в Excel без повторного обращения к API.

Запуск отдельно: python xlsx_export.py output/document.json
Зависимость приложения: openpyxl>=3.1.5,<4.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import tempfile
from pathlib import Path


def _scalar(value):
    """Сохранять исходные типы; вложенные значения представлять как JSON."""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def _flatten(node, prefix=""):
    """Развернуть реквизиты, отделив значение от служебной информации поля."""
    if isinstance(node, dict) and "value" in node:
        yield prefix, node
    elif isinstance(node, dict):
        for key, value in node.items():
            yield from _flatten(value, f"{prefix}.{key}" if prefix else key)
    else:
        yield prefix, {"value": node}


def export_result_xlsx(
    result: dict,
    output_path: str | Path,
    *,
    include_raw_text: bool | None = None,
    include_details: bool | None = None,
) -> Path:
    """Создать Excel с товарами, реквизитами и итогами.

    Исходный текст и детализация по умолчанию не включаются в Excel.
    Вернуть их можно параметрами include_raw_text=True / include_details=True
    или переменными XLSX_INCLUDE_RAW_TEXT=true / XLSX_INCLUDE_DETAILS=true.
    Полный исходный JSON при этом не изменяется.
    declaration_draft остаётся только во внутреннем JSON и не экспортируется.

    Значения не пересчитываются и не исправляются. Пустые значения остаются
    пустыми; строки, похожие на формулы, сохраняются именно текстом.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    except ImportError as error:
        raise RuntimeError(
            "Для Excel установите openpyxl: %pip install 'openpyxl>=3.1.5,<4'. "
            "JSON уже сохранён; его можно экспортировать без нового запроса к ИИ."
        ) from error

    if not isinstance(result, dict):
        raise TypeError("Ожидается JSON-объект результата модели")
    # Отдельное представление для Excel: исходный ответ и сохранённый JSON
    # остаются полными. Черновик исключён также из необязательной детализации.
    result = {key: value for key, value in result.items() if key != "declaration_draft"}
    if include_raw_text is None:
        include_raw_text = os.getenv("XLSX_INCLUDE_RAW_TEXT", "false").lower() in {"true", "1", "yes"}
    if include_details is None:
        include_details = os.getenv("XLSX_INCLUDE_DETAILS", "false").lower() in {"true", "1", "yes"}
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(name, headers, rows):
        ws = workbook.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append([_scalar(value) for value in row])
        for row in ws:
            for cell in row:
                # OCR-текст не должен превращаться в исполняемые Excel-формулы.
                if isinstance(cell.value, str):
                    if ILLEGAL_CHARACTERS_RE.search(cell.value):
                        raise ValueError(f"Недопустимый XML-символ в {name}!{cell.coordinate}")
                    if len(cell.value) > 32767:
                        raise ValueError(f"Превышен лимит Excel 32767 символов: {name}!{cell.coordinate}")
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="193B58")
            cell.font = Font(color="FFFFFF", bold=True)
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            width = max(len(str(cell.value or "")) for cell in column)
            ws.column_dimensions[get_column_letter(column[0].column)].width = min(55, max(16, width + 2))
        return ws

    # В Excel выводим только данные исходных документов.
    scopes = []
    documents = result.get("documents", [])
    if isinstance(documents, dict):
        documents = [documents]
    for index, document in enumerate(documents or [], 1):
        if not isinstance(document, dict):
            raise TypeError("Элемент documents должен быть объектом")
        scopes.append(("Документ", str(document.get("document_id", index)), document))
    if not scopes:
        scopes.append(("Результат", "", result))

    # Одна товарная позиция = одна строка. Поля берём из фактического JSON,
    # включая новые лейблы; единицы измерения выводим соседними столбцами.
    item_records = []
    field_names = []
    for scope, document_id, document in scopes:
        for index, item in enumerate(document.get("items", []) or [], 1):
            fields = dict(_flatten(item))
            for name in fields:
                if name not in field_names:
                    field_names.append(name)
            item_records.append((scope, document_id, index, fields))
    item_rows = []
    for scope, document_id, index, fields in item_records:
        row = [scope, document_id, index]
        for name in field_names:
            field = fields.get(name, {})
            row.extend([field.get("value"), field.get("unit")])
        item_rows.append(row)
    sheet("Товары", ["Область", "Документ", "Позиция"] + [
        column for name in field_names for column in (name, f"{name}: единица")
    ], item_rows)

    # Реквизиты и итоговые значения выводим отдельно от товарных строк.
    requisites = []
    totals = []
    for scope, document_id, document in scopes:
        for key, value in document.items():
            if key == "items":
                continue
            for name, field in _flatten(value, key):
                row = [scope, document_id, name, field.get("value"), field.get("unit")]
                if include_raw_text:
                    row.append(field.get("raw_text"))
                (totals if key == "totals" else requisites).append(row)
    headers = ["Область", "Документ", "Поле", "Значение", "Единица"]
    if include_raw_text:
        headers.append("Исходный текст")
    # Единицы сохраняются в товарах и итогах; в реквизитах столбец не нужен.
    sheet("Реквизиты", headers[:4] + headers[5:], [row[:4] + row[5:] for row in requisites])
    sheet("Итоги", headers, totals)

    # Полный ответ по JSON-путям: сохраняются confidence, evidence, bbox,
    # проверки, конфликты и прочие ключи, даже не известные экспортёру заранее.
    details = []

    def walk(value, path="$"):
        if isinstance(value, dict) and value:
            for key, child in value.items():
                escaped = str(key).replace("~", "~0").replace("/", "~1")
                walk(child, f"{path}/{escaped}")
        elif isinstance(value, list) and value:
            for index, child in enumerate(value):
                walk(child, f"{path}/{index}")
        else:
            details.append([path, "null" if value is None else type(value).__name__, value])

    if include_details:
        walk(result)
        sheet("Детализация", ["JSON-путь", "Тип", "Значение"], details)
    # Атомарная запись: при ошибке готовый файл предыдущего запуска не портится.
    handle, temporary_name = tempfile.mkstemp(suffix=".xlsx", dir=output_path.parent)
    os.close(handle)
    try:
        workbook.save(temporary_name)
        os.replace(temporary_name, output_path)
    finally:
        workbook.close()
        Path(temporary_name).unlink(missing_ok=True)
    return output_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("json_path", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    data = json.loads(args.json_path.read_text(encoding="utf-8"))
    print(export_result_xlsx(data, args.output or args.json_path.with_suffix(".xlsx")))
